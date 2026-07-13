"""
Core utility functions — text processing, PDF generation, resume helpers.

NOTE: This module is the single source of truth for text processing functions.
Do NOT duplicate these in ml_service.py or elsewhere.
"""
from __future__ import annotations
import re
import io
from datetime import datetime

from fastapi import UploadFile
from app.core.roles import JOB_ROLES


# ═══════════════════════════════════════════════════════════════════════════════
# Text Processing
# ═══════════════════════════════════════════════════════════════════════════════

def remove_emojis(text: str) -> str:
    """Remove non-latin-1 characters (emojis, special Unicode)."""
    return text.encode('latin-1', 'ignore').decode('latin-1').strip()


def clean_resume(text: str) -> str:
    """Normalize resume text: remove URLs, special chars, lowercase."""
    if not isinstance(text, str):
        return ""
    text = re.sub(r'http\S+\s*', ' ', text)
    text = re.sub(r'RT|cc', ' ', text)
    text = re.sub(r'#\S+', '', text)
    text = re.sub(r'@\S+', '  ', text)
    text = re.sub(r'[%s]' % re.escape("""!"#$%&'()*+,-./:;<=>?@[\]^_`{|}~"""), ' ', text)
    text = re.sub(r'[^\x00-\x7f]', r' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.lower().strip()


# ═══════════════════════════════════════════════════════════════════════════════
# File Extraction (legacy — prefer file_service.py for new code)
# ═══════════════════════════════════════════════════════════════════════════════

async def extract_text(file: UploadFile) -> str:
    """Extract raw text from uploaded PDF/DOCX/TXT file."""
    try:
        content = await file.read()
        await file.seek(0)
        filename = (file.filename or "").lower()
        if filename.endswith('.pdf'):
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return " ".join(page.extract_text() or "" for page in reader.pages).strip()
        elif filename.endswith('.docx'):
            from docx import Document
            doc = Document(io.BytesIO(content))
            return " ".join(para.text for para in doc.paragraphs if para.text.strip())
        elif filename.endswith('.txt'):
            return content.decode('utf-8')
        return ""
    except Exception:
        return ""


# ═══════════════════════════════════════════════════════════════════════════════
# Role Helpers
# ═══════════════════════════════════════════════════════════════════════════════

def get_all_job_roles() -> tuple[dict, dict]:
    """Return (ui_roles_with_emoji, pdf_safe_roles) tuple."""
    pdf_safe = JOB_ROLES
    ui_roles = {}
    for k, v in pdf_safe.items():
        if "ML" in k or "AI" in k:
            ui_roles[f"🤖 {k}"] = v
        elif "Data" in k:
            ui_roles[f"🔬 {k}"] = v
        else:
            ui_roles[f"💻 {k}"] = v
    return ui_roles, pdf_safe


# ═══════════════════════════════════════════════════════════════════════════════
# PDF Report Generation (legacy — prefer report_service.py for new code)
# ═══════════════════════════════════════════════════════════════════════════════

def create_pdf_report(results: dict, role_data: dict | None = None) -> bytes:
    """Generate a PDF report from analysis results."""
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 20)
            self.cell(0, 10, 'AI Talent Scout Pro - Analysis Report', 0, 1, 'C')
            self.ln(10)

        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 10)
            self.cell(0, 10, f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}', 0, 0, 'C')

    pdf = PDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    clean_role = remove_emojis(results.get("role", "Unknown"))
    pdf.cell(0, 10, f'Role Analyzed: {clean_role}', 0, 1, 'C')
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'ANALYSIS SCORES:', 0, 1)
    pdf.set_font('Arial', '', 11)
    pdf.cell(0, 10, f'Overall Fit:        {int(results.get("role_fit", 0))}%', 0, 1)
    pdf.cell(0, 10, f'ML Proficiency:     {int(results.get("ml_score", 0))}/100', 0, 1)
    pdf.cell(0, 10, f'Semantic Score:     {int(results.get("semantic_score", 0))}%', 0, 1)
    pdf.cell(0, 10, f'Skill Similarity:   {int(results.get("skill_similarity", 0))}%', 0, 1)
    pdf.cell(0, 10, f'ATS Score:          {int(results.get("ats_score", 0))}/100', 0, 1)
    pdf.cell(0, 10, f'Experience:         {results.get("experience", 0)} years', 0, 1)
    pdf.ln(5)

    decision = results.get("decision", "PENDING")
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, f'RECOMMENDATION: {decision}', 0, 1, 'C')
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'TOP RECOMMENDED ROLES:', 0, 1)
    pdf.set_font('Arial', '', 10)
    for i, (role, score) in enumerate(results.get('top_roles', [])[:5], 1):
        clean_r = remove_emojis(str(role))
        pdf.cell(0, 8, f'{i}. {clean_r} ({int(score)}%)', 0, 1)

    pdf_buffer = io.BytesIO()
    pdf.output(pdf_buffer)
    pdf_buffer.seek(0)
    return pdf_buffer.getvalue()